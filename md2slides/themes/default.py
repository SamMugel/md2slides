"""Default theme for md2slides."""

from pptx.dml.color import RGBColor
from pptx.util import Pt

from md2slides.themes.base import BaseTheme, ThemeColors, ThemeFonts


class DefaultTheme(BaseTheme):
    """Default theme with a clean, professional look."""

    name = "default"

    colors = ThemeColors(
        background=(255, 255, 255),       # White
        title_text=(0, 51, 102),          # Dark blue
        body_text=(51, 51, 51),           # Dark gray
        accent=(0, 112, 192),             # Blue accent
        code_background=(245, 245, 245),  # Light gray
    )

    fonts = ThemeFonts(
        title_font="Calibri",
        body_font="Calibri",
        code_font="Courier New",
        title_size=44,
        subtitle_size=32,
        body_size=18,
        code_size=12,
    )

    def apply_to_title_slide(self, slide):
        """Apply theme settings to a title slide."""
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = self.fonts.title_font
                        run.font.color.rgb = RGBColor(*self.colors.title_text)

    def apply_to_content_slide(self, slide):
        """Apply theme settings to a content slide."""
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue

            is_title = shape == slide.shapes.title

            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    # Skip code runs (already styled with monospace font)
                    if run.font.name == self.fonts.code_font:
                        continue

                    if is_title:
                        run.font.name = self.fonts.title_font
                        run.font.size = Pt(self.fonts.title_size)
                        run.font.color.rgb = RGBColor(*self.colors.title_text)
                    else:
                        run.font.name = self.fonts.body_font
                        run.font.size = Pt(self.fonts.body_size)
                        run.font.color.rgb = RGBColor(*self.colors.body_text)
