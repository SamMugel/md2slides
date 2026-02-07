"""Light theme for md2slides."""

from pptx.dml.color import RGBColor
from pptx.util import Pt

from md2slides.themes.base import BaseTheme, ThemeColors, ThemeFonts


class LightTheme(BaseTheme):
    """Light theme with white background and minimal colors."""

    name = "light"

    colors = ThemeColors(
        background=(255, 255, 255),       # Pure white
        title_text=(50, 50, 50),          # Near black
        body_text=(80, 80, 80),           # Dark gray
        accent=(100, 100, 100),           # Medium gray
        code_background=(250, 250, 250),  # Very light gray
    )

    fonts = ThemeFonts(
        title_font="Arial",
        body_font="Arial",
        code_font="Consolas",
        title_size=44,
        subtitle_size=32,
        body_size=18,
        code_size=12,
    )

    def apply_to_title_slide(self, slide):
        """Apply theme settings to a title slide."""
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = self.fonts.title_font
                        run.font.color.rgb = RGBColor(*self.colors.title_text)

    def apply_to_content_slide(self, slide):
        """Apply theme settings to a content slide."""
        # Red color used for H3 inline headings - should not be overridden
        H3_RED = RGBColor(0xFF, 0x00, 0x00)

        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue

            is_title = shape == slide.shapes.title

            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    # Skip code runs (already styled with monospace font)
                    if run.font.name == self.fonts.code_font:
                        continue

                    # Skip H3 inline headings (styled with red color)
                    try:
                        if run.font.color.rgb == H3_RED:
                            continue
                    except (TypeError, AttributeError):
                        pass

                    if is_title:
                        run.font.name = self.fonts.title_font
                        run.font.size = Pt(self.fonts.title_size)
                        run.font.color.rgb = RGBColor(*self.colors.title_text)
                    else:
                        run.font.name = self.fonts.body_font
                        run.font.size = Pt(self.fonts.body_size)
                        run.font.color.rgb = RGBColor(*self.colors.body_text)
